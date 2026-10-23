"""Build the evidence-led defense report from its editable Markdown source.

Requires the bundled python-docx and Pillow. It does not read private configuration.
Render and visually inspect the result with the documents skill before delivery.
"""
from pathlib import Path
from xml.sax.saxutils import escape
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/defense'
ASSETS = OUT / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
FONT_PATH = Path('C:/Windows/Fonts/msyh.ttc')
if not FONT_PATH.exists():
    raise SystemExit('Set FONT_PATH to an installed CJK font before building the report.')


class Diagram:
    def __init__(self, width, height):
        self.width, self.height = width, height
        self.image = Image.new('RGB', (width, height), 'white')
        self.draw = ImageDraw.Draw(self.image)
        self.svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
                    '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="9" refY="5" orient="auto"><path d="M0,0 L10,5 L0,10" fill="#566579"/></marker></defs>',
                    '<rect width="100%" height="100%" fill="white"/>']

    def text(self, xy, lines, size=29, color='#152234'):
        if isinstance(lines, str):
            lines = lines.split('\n')
        font = ImageFont.truetype(str(FONT_PATH), size)
        x, y = xy
        step = round(size * 1.45)
        start = y - (len(lines)-1) * step / 2
        for index, line in enumerate(lines):
            top = start + index * step
            self.draw.text((x, top), line, fill=color, font=font, anchor='mm')
            self.svg.append(f'<text x="{x}" y="{top}" text-anchor="middle" dominant-baseline="central" font-family="Microsoft YaHei,sans-serif" font-size="{size}" fill="{color}">{escape(line)}</text>')

    def box(self, bounds, lines, size=29, fill='#edf2f7'):
        x1, y1, x2, y2 = bounds
        self.draw.rounded_rectangle(bounds, radius=12, fill=fill, outline='#718198', width=2)
        self.svg.append(f'<rect x="{x1}" y="{y1}" width="{x2-x1}" height="{y2-y1}" rx="12" fill="{fill}" stroke="#718198" stroke-width="2"/>')
        self.text(((x1+x2)/2, (y1+y2)/2), lines, size)

    def arrow(self, points, label=None, label_xy=None):
        self.draw.line(points, fill='#566579', width=3)
        import math
        x, y = points[-1]
        before_x, before_y = points[-2]
        angle = math.atan2(y-before_y, x-before_x)
        triangle = [(x,y), (x-16*math.cos(angle-.4), y-16*math.sin(angle-.4)),
                    (x-16*math.cos(angle+.4), y-16*math.sin(angle+.4))]
        self.draw.polygon(triangle, fill='#566579')
        path = ' '.join(f'{a},{b}' for a,b in points)
        self.svg.append(f'<polyline points="{path}" fill="none" stroke="#566579" stroke-width="3" marker-end="url(#arrow)"/>')
        if label:
            self.text(label_xy, label, size=23)

    def save(self, name):
        self.image.save(ASSETS / (name+'.png'), dpi=(240,240))
        (ASSETS / (name+'.svg')).write_text('\n'.join(self.svg+['</svg>']), encoding='utf-8')


def diagrams():
    a = Diagram(1600, 900)
    for bounds, label in [((60,30,500,155),'Vue 用户端'),((580,30,1020,155),'Vue 管理后台'),((1100,30,1540,155),'微信原生小程序')]:
        a.box(bounds,label,34)
        x=(bounds[0]+bounds[2])/2
        a.arrow([(x,155),(x,225),(800,225),(800,270)])
    a.text((1130,198),'统一 HTTP 接口与 Bearer 会话',26)
    a.box((160,270,1440,370),'Spring Boot 服务端   Java 17',36)
    a.box((160,410,740,545),['身份与对象权限','账号  认证  文件  角色'],31)
    a.box((860,410,1440,545),['业务服务与数据库事务','商品  任务  订单  私信  通知'],31)
    a.arrow([(450,370),(450,410)])
    a.arrow([(1150,370),(1150,410)])
    a.box((160,625,740,755),['MySQL 业务数据与持久任务','状态  账本  审计  幂等结果'],31)
    a.box((860,625,1440,755),['受控文件与外部能力适配','扫描  微信换码  后续支付短信'],31)
    a.arrow([(450,545),(450,625)])
    a.arrow([(1010,545),(1010,580),(580,580),(580,625)])
    a.arrow([(1150,545),(1150,625)])
    a.text((800,838),'支付短信未接通   真实微信与正式部署仍待验收',28)
    a.save('architecture')

    e = Diagram(1600, 1030)
    e.box((580,30,1020,145),['用户 users','发布者  买家  接取者'],30)
    e.box((60,255,470,380),['商品 product_info','分类  价格  库存'],30)
    e.box((600,255,1040,380),['任务 RewardTasks','奖励  状态  校区'],30)
    e.box((1150,255,1550,380),['奖励账户 reward_accounts','可用与冻结金额'],28)
    e.arrow([(700,145),(700,200),(265,200),(265,255)],'1 对多',(390,179))
    e.arrow([(820,145),(820,255)],'1 对多',(910,201))
    e.arrow([(930,145),(930,200),(1350,200),(1350,255)],'按奖励类型',(1225,178))
    e.box((60,505,470,635),['订单 user_orders','参与方  价格快照  状态'],29)
    e.box((600,505,1040,635),['申请 task_applications','凭证 task_proofs'],30)
    e.box((1150,505,1550,635),['流水 reward_ledger','任务事件键唯一'],29)
    e.arrow([(265,380),(265,505)],'1 对多',(360,445))
    e.arrow([(710,380),(710,440),(470,440),(470,505)])
    e.arrow([(820,380),(820,505)],'1 对多',(910,445))
    e.arrow([(1350,380),(1350,505)],'1 对多',(1440,445))
    e.box((600,795,1040,925),['结算 task_settlement','每个任务最多一次成功结算'],29)
    e.box((60,795,470,925),['争议 order_aftersales','申请时订单上下文'],29)
    e.arrow([(265,635),(265,795)],'1 对多',(365,718))
    e.arrow([(1040,325),(1080,325),(1080,760),(820,760),(820,795)],'每任务至多一份',(995,714))
    e.arrow([(1150,570),(1095,570),(1095,860),(1040,860)])
    e.text((800,990),'局部关系示意   完整外键与索引以结构和迁移为准',25)
    e.save('core-relations')

    f = Diagram(1600, 900)
    f.box((70,60,390,180),['发布并冻结','待审核'],31)
    f.box((640,60,960,180),['审核通过','待接取'],31)
    f.box((1210,60,1530,180),['分配接取者','进行中'],31)
    f.arrow([(390,120),(640,120)])
    f.arrow([(960,120),(1210,120)],'申请待审再分配',(1085,83))
    f.box((1210,390,1530,510),['提交最新凭证','待验收'],31)
    f.arrow([(1370,180),(1370,390)])
    f.box((640,390,960,510),['验收通过','已完成并唯一结算'],29)
    f.arrow([(1210,450),(960,450)])
    f.arrow([(1530,450),(1570,450),(1570,235),(1450,235),(1450,180)],'补证',(1507,278))
    f.box((70,390,390,510),['审核驳回','已驳回并释放冻结'],29)
    f.arrow([(230,180),(230,390)],'驳回',(305,280))
    f.box((640,690,960,810),['允许关闭时管理关闭','已取消并释放冻结'],28)
    f.arrow([(800,180),(800,270),(515,270),(515,750),(640,750)],'符合关闭条件',(580,610))
    f.box((1210,690,1530,810),['凭证未通过','未完成且继续冻结'],29)
    f.arrow([(1370,510),(1370,690)])
    f.arrow([(1210,750),(960,750)])
    f.text((800,862),'进行中与待验收不能直接管理关闭   申请不等于已经接取',25)
    f.save('task-flow')


def east_asia(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0,0,0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0,fonts)
    fonts.set(qn('w:eastAsia'), '宋体' if not bold else '黑体')
    color = rpr.find(qn('w:color'))
    if color is not None:
        for key in ('themeColor','themeTint','themeShade'):
            color.attrib.pop(qn('w:'+key), None)
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        for tag in ('w:pBdr','w:shd'):
            for node in list(ppr.findall(qn(tag))):
                ppr.remove(node)


def inline(paragraph, text):
    for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text):
        if not part:
            continue
        if part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1]); run.font.name='Consolas'
            run.font.size=Pt(9.5)
        elif part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold=True
        else:
            paragraph.add_run(part)


def table(doc, rows):
    t=doc.add_table(rows=1,cols=len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    count=len(rows[0])
    widths={2:[1.8,4.7],3:[1.25,2.68,2.57]}.get(count,[6.5/count]*count)
    if count==3 and rows[0][0]=='编号':
        widths=[.45,3.5,2.55]
    if count==3 and rows[0][0]=='示例路径':
        widths=[2.5,1.2,2.8]
    for col, width in zip(t.columns,widths):
        col.width=Inches(width)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
        edge=OxmlElement('w:'+name)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:
            edge.set(qn('w:'+key),value)
        borders.append(edge)
    props.append(borders)
    header=t.rows[0]._tr.get_or_add_trPr()
    header.append(OxmlElement('w:tblHeader'))
    for index, values in enumerate(rows):
        cells=t.rows[0].cells if index==0 else t.add_row().cells
        for column,(cell,value) in enumerate(zip(cells,values)):
            cell.width=Inches(widths[column])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for name,size in [('top',95),('bottom',95),('left',110),('right',110)]:
                node=OxmlElement('w:'+name);node.set(qn('w:w'),str(size));node.set(qn('w:type'),'dxa');margins.append(node)
            tcpr.append(margins)
            shading=OxmlElement('w:shd')
            shading.set(qn('w:fill'),'E7EDF4' if index==0 else ('F7F9FB' if index%2==0 else 'FFFFFF'))
            tcpr.append(shading)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.2
            p.paragraph_format.keep_with_next=False
            p.paragraph_format.widow_control=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if index==0 or (column==0 and len(value)<16) else WD_ALIGN_PARAGRAPH.LEFT
            inline(p,value)
            for run in p.runs:
                run.font.size=Pt(10)
                run.bold=index==0
            trpr=cell._tc.getparent().get_or_add_trPr()
            if trpr.find(qn('w:cantSplit')) is None:
                trpr.append(OxmlElement('w:cantSplit'))
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


def build():
    diagrams()
    source=(OUT/'校园Π软件工程设计报告答辩版.md').read_text(encoding='utf-8').splitlines()
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.82);sec.bottom_margin=Inches(.78)
    sec.left_margin=Inches(1);sec.right_margin=Inches(1)
    sec.header_distance=Inches(.33);sec.footer_distance=Inches(.33)
    east_asia(doc.styles['Normal'],11.5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.35
    doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.widow_control=True
    for name,size in [('Title',25),('Subtitle',14),('Heading 1',17),('Heading 2',13)]:
        east_asia(doc.styles[name],size,bold=name not in ('Subtitle',))
        doc.styles[name].paragraph_format.keep_with_next=True
        doc.styles[name].paragraph_format.space_before=Pt(15 if name.startswith('Heading') else 0)
        doc.styles[name].paragraph_format.space_after=Pt(9)
    sec.different_first_page_header_footer=True
    hp=sec.header.paragraphs[0]
    hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=hp.add_run('校园Π软件工程设计报告  答辩版');r.font.size=Pt(9);r.font.color.rgb=RGBColor(0,0,0)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=fp.add_run('第 ');r.font.size=Pt(9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
    fp.add_run(' 页').font.size=Pt(9)
    lines=iter(enumerate(source))
    in_cover=True
    cover_intro=False
    for _,line in lines:
        line=line.strip()
        if not line:
            continue
        if line.startswith('# '):
            doc.add_paragraph().paragraph_format.space_after=Pt(54)
            title=line[2:].replace('软件工程设计报告','\n软件工程设计报告')
            p=doc.add_paragraph(title,style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        elif line=='## 摘要':
            doc.add_page_break();in_cover=False
            doc.add_heading('摘要',level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:],level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:],level=2)
        elif line.startswith('!['):
            match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
            if not match:
                raise ValueError('Invalid image line')
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(OUT/match[2]),width=Inches(6.45))
        elif line.startswith('|'):
            rows=[]
            while True:
                cells=[x.strip() for x in line.strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v) for v in cells):
                    rows.append(cells)
                try:
                    _,line=next(lines);line=line.strip()
                except StopIteration:
                    line='';break
                if not line.startswith('|'):
                    break
            table(doc,rows)
            if line:
                raise ValueError('Tables must be separated by a blank line')
        elif line.startswith('图 '):
            p=doc.add_paragraph(line);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(12)
            for run in p.runs:run.font.size=Pt(10)
        else:
            p=doc.add_paragraph()
            inline(p,line)
            if in_cover:
                if line.startswith('本报告依据'):
                    p.paragraph_format.space_before=Pt(20)
                    cover_intro=True
                elif not cover_intro:
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after=Pt(14)
                    if line=='软件工程项目设计与实现答辩材料':
                        p.style='Subtitle';p.paragraph_format.space_after=Pt(38)
    # All titles/headings must remain black; remove inherited theme decoration.
    for p in doc.paragraphs:
        if p.style.name in ('Title','Subtitle','Heading 1','Heading 2'):
            ppr=p._p.get_or_add_pPr()
            for tag in ('w:pBdr','w:shd'):
                for node in list(ppr.findall(qn(tag))):ppr.remove(node)
            for run in p.runs:
                run.font.color.rgb=RGBColor(0,0,0)
                run.font.underline=False
    doc.core_properties.title='校园Π校园信息共享平台软件工程设计报告'
    doc.core_properties.subject='依据实际实现和验证记录编写的答辩材料'
    doc.core_properties.author='校园Π项目组'
    doc.core_properties.keywords='校园Π 软件工程 任务结算 组件复用 答辩'
    target=OUT/'校园Π软件工程设计报告答辩版.docx'
    doc.save(target)
    print('Created',target,'paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'figures',len(doc.inline_shapes))


if __name__=='__main__':
    build()
